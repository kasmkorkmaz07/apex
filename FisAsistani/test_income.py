import unittest,tempfile
from pathlib import Path
from copy import deepcopy
from openpyxl import load_workbook
from core import draft,validate,money,vat_records,Store,gross_line
from accounting import profile,save_profile,EMIR,resolve

SAMPLES=[
('184','25.07.2026',{'1':'350'},'''ESER SİMİT FIRINI
TARİH: 25-07-2026
FİŞ NO: 0006
Z GÜNLÜK RAPORU
Z SAYAÇ 0184
SİMİT %1 5
*350,00
NET SATIŞ *350,00
VERGİ DÖKÜMÜ
%1 TOPLAM *350,00
TOPKDV *3,46
TOPLAM *350,00
KÜM. TOPLAM *67.266,11
EKÜ NO:0001 Z NO:0184'''),
('3215','01.08.2026',{'10':'4950'},'''NAZLIM KOZMETİK
01-08-2026 FİŞ NO 16
Z RAPORU
Z NO: 3.215
TUHAFİYE 8
*4 950,00
FİŞ İPTAL *1.780,00
TOPLAM %10 *4 950,00
TOPKDV %10 *450,00
TOPLAM *4 950,00
TOPKDV *450,00
KÜM. TOPLAM *6.185.416,72
EKÜ NO: 1 Z NO 3.215'''),
('817','29.07.2026',{'0':'20774.75','10':'1905'},'''İKİ Bİ DÜNYA BİLGİ
29.07.2026 Saat 21:02
Fiş No: 0047
Z RAPORU
Z No: 0817 EKÜ No:0003
KİTAP %0 Mktr 39,00
*20.774,75
KIRTASİYE-10 %10 Mktr 3,000
*1.905,00
TOPLAM %0 *20.774,75
KDV %0 *0,00
TOPLAM %10 *1.905,00
KDV %10 *173,18
TOPLAM SATIŞ *22.679,75
TOPLAM KDV *173,18
KÜMSATIŞ *29.138.701,11'''),
('796','05.02.2026',{'1':'5618'},'''ALİ KALENDER
05/02/2026 FİŞ NO:0010
Z GÜNLÜK RAPORU
TARİH:05/02/26
Z SAYAÇ 796
ET %01 6,000
*4.848,00
ÇİĞ TAVUK %01 3,000
*770,00
NET SATIŞ *5.618,00
VERGİ DÖKÜMÜ
% 01 TOPLAM *5.618,00
KDV *55,62
TOP *5.618,00
KDV *55,62
KÜM TOP *2.054.777,32
EKÜ NO:0001 Z NO:0796''')]

class IncomeTests(unittest.TestCase):
 def test_four_printed_formats(self):
  for no,date,rates,text in SAMPLES:
   with self.subTest(no=no):
    d=draft(text,'gelir');validate(d)
    self.assertEqual(d['evrak_no'],no);self.assertEqual(d['evrak_tarihi'],date)
    self.assertEqual(d['aciklama'],'Z RAPORU')
    self.assertEqual({str(int(money(r['oran']))):money(r['matrah'])+money(r['kdv']) for r in d['kalemler']},{k:money(v) for k,v in rates.items()})
    self.assertEqual(sum(money(x['toplam']) for x in vat_records(d)),money(d['toplam']))
 def test_ambiguous_z_never_uses_receipt_number(self):
  d=draft(SAMPLES[0][3]+'\nZ NO: 185','gelir')
  self.assertEqual(d['evrak_no'],'')
  with self.assertRaises(ValueError):validate(d)
 def test_total_mismatch_blocks(self):
  d=draft(SAMPLES[0][3].replace('TOPLAM *350,00\nKÜM','TOPLAM *360,00\nKÜM'),'gelir')
  with self.assertRaises(ValueError):validate(d)
 def test_business_zero_hz_date_and_description(self):
  with tempfile.TemporaryDirectory() as tmp:
   s=Store(Path(tmp)/'db');self.addCleanup(s.db.close)
   cid=s.client('Hizmet');p=profile(s,cid);p['income_code']='HZ';save_profile(s,cid,p)
   rid=s.add(cid,None,draft(SAMPLES[2][3],'gelir'),'08-2026')
   for i,d,_,_ in s.rows(cid):s.save(i,d,True)
   self.assertEqual(s.receipt_period(rid),'07-2026')
   self.assertEqual(s.monthly_records(cid,'08-2026'),[])
   path=Path(tmp)/'out.xlsx';s.sync_month(cid,'07-2026',path)
   w=load_workbook(path);ws=w['Gelir_Aktarim']
   self.assertEqual(ws.max_column,6);self.assertEqual(ws.max_row,2)
   self.assertTrue(all(r[0]=='HZ' and r[1].strftime('%d.%m.%Y')=='29.07.2026' and r[3]=='Z RAPORU' for r in ws.values))
   self.assertEqual(sum(r[5] for r in ws.values),22679.75)
   self.assertEqual(s.vendor_rows(),[]);w.close();s.db.close()
 def test_general_optional_id_and_separate_accounts(self):
  with tempfile.TemporaryDirectory() as tmp:
   s=Store(Path(tmp)/'db');self.addCleanup(s.db.close);cid=s.client('Genel')
   p=deepcopy(EMIR);p.update(income_closing='100.01',income_rates={'0':['600.00',''],'1':['600.01','391.01'],'10':['600.10','391.10']});save_profile(s,cid,p)
   for text in [SAMPLES[0][3],SAMPLES[2][3]]:s.add(cid,None,draft(text,'gelir'),'08-2026')
   d=dict(tarih='01.08.2026',evrak_tarihi='01.08.2026',evrak_no='G1',aciklama='GİDER PUSULASI',vkn='1234567890',kind='gider',kod='MS',toplam='110',kalemler=[gross_line(10,110)],ocr='')
   s.add(cid,None,d,'08-2026')
   for i,d,_,_ in s.rows(cid):s.save(i,d,True)
   path=Path(tmp)/'out.xlsx';s.sync_month(cid,'08-2026',path);w=load_workbook(path)
   self.assertEqual(w['Genel_Satis'].max_column,11);self.assertEqual(w['Genel_Alis'].max_column,10)
   self.assertIsNone(w['Genel_Satis_Aktarim']['A1'].value)
   self.assertEqual(w['Genel_Alis_Aktarim']['J1'].value,'320 01')
   w.close();s.sync_month(cid,'07-2026',path);w=load_workbook(path)
   self.assertEqual(w['Genel_Satis_Aktarim'].max_row,3)
   for r in w['Genel_Satis_Aktarim'].values:
    self.assertEqual(r[9],'100.01');self.assertTrue(r[7].startswith('600.'));self.assertEqual(r[10],'Zraporu');self.assertIsNone(r[3])
   w.close();s.db.close()
 def test_missing_income_account_blocks(self):
  with self.assertRaises(ValueError):resolve({'kind':'gelir'},EMIR,10)
 def test_template_survives_recall(self):
  with tempfile.TemporaryDirectory() as tmp:
   s=Store(Path(tmp)/'db');s.add_vendor('1234567890','ESKİ ÜNVAN')
   d={'kind':'gelir','vkn':'1234567890','aciklama':'Z RAPORU'}
   self.assertEqual(s.recall_vendor(d)['aciklama'],'Z RAPORU');s.db.close()
 def test_ocr_punctuation_and_cross_checked_rate(self):
  d=draft(SAMPLES[1][3].replace('TOPLAM %10','TOPLAM %.10'),'gelir');validate(d)
  self.assertEqual(d['kalemler'][0]['oran'],'10.00')
  d=draft(SAMPLES[2][3].replace('TOPLAM %10','TOPLAM %19'),'gelir');validate(d)
  self.assertEqual(d['kalemler'][1]['oran'],'10.00')
  from review import warnings
  self.assertTrue(warnings(d))
 def test_general_mixed_ids_keeps_column(self):
  from accounting import export
  with tempfile.TemporaryDirectory() as tmp:
   p=deepcopy(EMIR);p.update(income_closing='100.01',income_rates={'1':['600.01','391.01']})
   a=draft(SAMPLES[0][3],'gelir');b=dict(a,evrak_no='185',vkn='1234567890',is_invoice=True,is_z_report=False)
   target=Path(tmp)/'out.xlsx';export(target,[(1,a),(2,b)],p)
   w=load_workbook(target);ws=w['Genel_Satis_Aktarim']
   self.assertEqual(ws.max_column,11);self.assertIsNone(ws['D1'].value);self.assertEqual(ws['D2'].value,'1234567890');w.close()
   w=load_workbook(target)
   try:
    self.assertNotIn('Genel_Z_Aktarim',w.sheetnames);self.assertNotIn('Genel_Fatura_Aktarim',w.sheetnames)
    self.assertEqual(w['Genel_Satis_Aktarim']['K1'].value,'Zraporu')
    self.assertEqual(w['Genel_Satis_Aktarim']['K2'].value,'Satış Faturası')
   finally:w.close()
