"""Mükellefe özel gelir/gider hesapları ve Zirve aktarımı."""
import json
from core import money, date_value, validate, validate_export
HEADERS=['Evrak Tarihi','Evrak No','Açıklama','VKN / TCKN','KDV Oranı','Matrah','KDV Tutarı','Matrah Hesap Kodu','KDV Hesap Kodu','Kapatılacak Hesap Kodu']
DESCRIPTIONS=['Z RAPORU','PERAKENDE SATIŞ FİŞİ','GİDER PUSULASI']
EMIR={'mode':'genel','closing':'320 01','rates':{'1':['740 01','191 01'],'10':['740 10','191 10'],'20':['740 20','191 20']}}

def profile(store,cid):
 row=store.db.execute('SELECT data FROM accounting_profiles WHERE client=?',(cid,)).fetchone()
 p=json.loads(row[0]) if row else {'mode':'isletme','closing':'','rates':{}}
 for k,v in {'income_code':'ML','income_closing':'','income_rates':{},'descriptions':DESCRIPTIONS,'income_description':'Z RAPORU'}.items():p.setdefault(k,v)
 return p

def save_profile(store,cid,p):
 if p['mode'] not in ('isletme','genel'):raise ValueError('Defter türünü seçin.')
 if p.get('income_code','ML') not in ('ML','HZ'):raise ValueError('Gelir kodu ML veya HZ olmalı.')
 for field in ('rates','income_rates'):
  for rate,codes in p.get(field,{}).items():
   if len(codes)!=2:raise ValueError('Her oran için iki hesap alanı gerekir.')
   if any(codes) and (not codes[0].strip() or (rate!='0' and not codes[1].strip())):raise ValueError('%'+rate+' için matrah ve KDV hesaplarını birlikte doldurun.')
 p['descriptions']=list(dict.fromkeys(x.strip() for x in p.get('descriptions',DESCRIPTIONS) if x.strip()))
 with store.db:store.db.execute('INSERT OR REPLACE INTO accounting_profiles VALUES (?,?)',(cid,json.dumps(p,ensure_ascii=False)))

def apply_defaults(store,cid,d):
 p=profile(store,cid);d=dict(d);d['accounting_mode']=p['mode']
 if d.get('kind')=='gelir':
  d['kod']=p['income_code']
  if not d.get('is_invoice'):d['vkn']=''
  if d.get('is_z_report'):d['aciklama']='Z RAPORU'
 return d

def resolve(d,p,rate):
 override=d.get('account_override',{});key=str(int(money(rate))) if money(rate)==int(money(rate)) else str(money(rate))
 income=d.get('kind')=='gelir';defaults=p.get('income_rates' if income else 'rates',{}).get(key,['',''])
 codes=[override.get('base',''),override.get('vat','')] if override else defaults[:]
 closing=override.get('closing','') if override else p.get('income_closing' if income else 'closing','')
 if not codes[0] or (money(rate)!=0 and not codes[1]) or not closing:raise ValueError(('Gelir' if income else 'Gider')+' %'+key+' için hesapları tamamlayın (%0 KDV hesabı gerektirmez).')
 if money(rate)==0:codes[1]=''
 return codes,closing

def export(path,records,p):
 from openpyxl import Workbook
 from openpyxl.styles import Font,PatternFill
 from openpyxl.utils import get_column_letter
 wb=Workbook();wb.remove(wb.active)
 sections=[('gider','Genel_Alis'),('gelir','Genel_Satis')]
 for kind,name in sections:
  selected=[(i,dict(d,accounting_mode='genel')) for i,d in records if d.get('kind','gider')==kind]
  headers=list(HEADERS)
  if kind=='gelir':headers.append('Evrak Türü')
  ws=wb.create_sheet(name);ws.append(headers)
  for rid,d in selected:
   validate_export(d,rid)
   for line in d['kalemler']:
    rate=money(line['oran'])
    if rate==0 and kind!='gelir':continue
    codes,closing=resolve(d,p,rate)
    values=[date_value(d['evrak_tarihi'] if kind=='gelir' else d['tarih']),d['evrak_no'],d['aciklama'],d.get('vkn',''),float(rate),float(money(line['matrah'])),float(money(line['kdv'])),*codes,closing]
    if kind=='gelir':values.append('Zraporu' if d.get('is_z_report') else 'Satış Faturası' if d.get('is_invoice') else 'Perakende Satış Fişi')
    ws.append(values)
    for col,title in enumerate(headers,1):
     cell=ws.cell(ws.max_row,col)
     if title=='Evrak Tarihi':cell.number_format='dd.mm.yyyy'
     elif title in ('KDV Oranı','Matrah','KDV Tutarı'):cell.number_format='#,##0.00'
     else:cell.data_type='s';cell.number_format='@'
  for c in ws[1]:c.font=Font(bold=True,color='FFFFFF');c.fill=PatternFill('solid',fgColor='234B58')
  ws.freeze_panes='A2';ws.auto_filter.ref=ws.dimensions
  for i,title in enumerate(headers,1):ws.column_dimensions[get_column_letter(i)].width=38 if title=='Açıklama' else max(18,len(title)+3)
  raw=wb.copy_worksheet(ws);raw.title=name+'_Aktarim';raw.delete_rows(1);raw.auto_filter.ref=None;raw.freeze_panes=None
 guide=wb.create_sheet('Sutun_Eslestirme');guide.append(['Aktarım sayfası','Excel sütun no','Alan','Yön'])
 for kind,name in sections:
  for col in wb[name][1]:guide.append([name+'_Aktarim',col.column,col.value,'Matrah: Alacak / Kapatma: Borç' if kind=='gelir' else 'Matrah: Borç / Kapatma: Alacak'])
 for col,width in zip('ABCD',(28,20,32,43)):guide.column_dimensions[col].width=width
 from excel_save import save_workbook
 save_workbook(wb,path)

def settings(app,cid):
 import tkinter as tk
 from tkinter import ttk
 from in_app_panel import panel
 win=panel(app.root,'Mükellef • Gelir / gider ayarları',scroll=True)
 p=profile(app.store,cid);frame=ttk.Frame(win.body,padding=18);frame.pack(fill='both',expand=True)
 ttk.Label(frame,text='DEFTER VE GELİR AYARLARI',style='Title.TLabel').pack(anchor='w')
 top=ttk.Frame(frame);top.pack(fill='x',pady=12)
 mode=tk.StringVar(value='Genel muhasebe' if p['mode']=='genel' else 'İşletme defteri');code=tk.StringVar(value=p['income_code'])
 ttk.Label(top,text='Defter türü').grid(row=0,column=0,sticky='w')
 ttk.Combobox(top,textvariable=mode,values=['İşletme defteri','Genel muhasebe'],state='readonly',width=16).grid(row=0,column=1,padx=10)
 code_label=ttk.Label(top,text='Gelir kodu');code_label.grid(row=0,column=2)
 code_box=ttk.Combobox(top,textvariable=code,values=['ML','HZ'],state='readonly',width=8);code_box.grid(row=0,column=3,padx=10)
 code_hint=ttk.Label(top,text='ML: Mal satışı   HZ: Hizmet satışı');code_hint.grid(row=1,column=1,columnspan=3,sticky='w',pady=8)
 description=tk.StringVar(value=p['income_description']);templates=list(p['descriptions'])
 ttk.Label(top,text='Açıklama şablonu').grid(row=2,column=0)
 descbox=ttk.Combobox(top,textvariable=description,values=templates,width=31);descbox.grid(row=2,column=1,columnspan=2,sticky='ew')
 def add_template():
  from template_dialog import ask_template
  value=ask_template(win)
  if value:
   if value not in templates:templates.append(value)
   descbox['values']=templates;description.set(value)
   latest=profile(app.store,cid);latest['descriptions']=list(templates);save_profile(app.store,cid,latest)
 def delete_template():
  from template_dialog import ask_template
  value=ask_template(win,templates,delete=True)
  if value:
   templates.remove(value);descbox['values']=templates
   if description.get()==value:description.set('')
   latest=profile(app.store,cid);latest['descriptions']=list(templates);save_profile(app.store,cid,latest)
 template_buttons=ttk.Frame(top);template_buttons.grid(row=2,column=3)
 ttk.Button(template_buttons,text='Şablon ekle',command=add_template).pack(side='left')
 ttk.Button(template_buttons,text='Sil',command=delete_template).pack(side='left',padx=3)
 ttk.Label(frame,text='Z raporu: Z RAPORU. Fatura: gelirde alıcı, giderde satıcı ünvanı. Belge üzerinde şablon seçilebilir.').pack(anchor='w')
 notebook=ttk.Notebook(frame);notebook.pack(fill='both',expand=True,pady=12)
 fields={};closings={}
 for kind,label in [('gider','GİDER HESAPLARI'),('gelir','GELİR HESAPLARI')]:
  page=ttk.Frame(notebook,padding=15);notebook.add(page,text=label)
  income=kind=='gelir';rates=p['income_rates' if income else 'rates']
  for col,title in enumerate(['KDV %','Satış / matrah hesabı' if income else 'Gider / matrah hesabı','Hesaplanan KDV hesabı' if income else 'İndirilecek KDV hesabı']):ttk.Label(page,text=title).grid(row=0,column=col,pady=8)
  fields[kind]={}
  for row,rate in enumerate(('0','1','10','20') if income else ('1','10','20'),1):
   ttk.Label(page,text='%'+rate).grid(row=row,column=0,padx=12)
   fields[kind][rate]=[tk.StringVar(value=v) for v in rates.get(rate,['',''])]
   for col,var in enumerate(fields[kind][rate],1):ttk.Entry(page,textvariable=var,state='disabled' if rate=='0' and col==2 else 'normal').grid(row=row,column=col,padx=5,pady=6,sticky='ew')
  closings[kind]=tk.StringVar(value=p['income_closing' if income else 'closing'])
  ttk.Label(page,text='Kapatma (borç)' if income else 'Kapatma (alacak)').grid(row=5,column=0,pady=14)
  ttk.Entry(page,textvariable=closings[kind]).grid(row=5,column=1,columnspan=2,sticky='ew')
  ttk.Label(page,text='Kullanmadığınız oranları boş bırakabilirsiniz. Kodları kendi Zirve hesap planınızdan girin.',wraplength=660).grid(row=6,column=0,columnspan=3,sticky='w')
  for col in (1,2):page.columnconfigure(col,weight=1)
 status=tk.StringVar(value='İşletme defterinde hesap tablosu gerekmez; gelir için ML / HZ seçimi yeterlidir.')
 status_label=ttk.Label(frame,textvariable=status,wraplength=740);status_label.pack(fill='x',pady=5)
 def save():
  try:
   updated=dict(p,mode='genel' if mode.get()=='Genel muhasebe' else 'isletme',income_code=code.get(),income_description=description.get().strip() or 'Z RAPORU',descriptions=templates)
   for kind,prefix in [('gider',''),('gelir','income_')]:
    updated[prefix+'rates']={r:[v.get().strip() for v in vs] for r,vs in fields[kind].items()};updated[prefix+'closing']=closings[kind].get().strip()
   save_profile(app.store,cid,updated);status.set('Kaydedildi. Excel çıkar / güncelle ile yeni ayarları uygulayın.')
  except ValueError as e:status.set(str(e))
 ttk.Button(frame,text='Ayarları kaydet',command=save,style='Accent.TButton').pack(anchor='e')
 def visibility(*_):
  general=mode.get()=='Genel muhasebe'
  for widget in (code_label,code_box,code_hint):
   if general:widget.grid_remove()
   else:widget.grid()
  if general:
   notebook.pack(fill='both',expand=True,pady=12,before=status_label)
   status.set('Gelir ve gider hesaplarını kendi Zirve hesap planınıza göre doldurun.')
  else:
   notebook.pack_forget()
   status.set('İşletme defteri için gelir kodunu ML veya HZ olarak seçin.')
 mode.trace_add('write',visibility);visibility()
 win.mode_var=mode;win.code_box=code_box;win.accounts=notebook;win.save_settings=save
 return win
