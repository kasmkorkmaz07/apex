"""İşletme defteri taslakları. Tutarlar Decimal, kimlikler metindir."""
from datetime import datetime
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP
import hashlib
import json
from pathlib import Path
import re
import sqlite3
import unicodedata

CENT = Decimal('0.01')

def turkish_sort(name):
    alphabet='ABCÇDEFGĞHIİJKLMNOÖPQRSŞTUÜVWXYZ'
    return tuple(alphabet.find(c) if c in alphabet else 100+ord(c) for c in name.replace('i','İ').replace('ı','I').upper())

def money(value):
    s = str(value).strip().replace(' ', '')
    if not s:
        raise ValueError('Tutar boş bırakılamaz.')
    if ',' in s:
        if not re.fullmatch(r'\d{1,3}(?:\.\d{3})*,\d{1,2}|\d+,\d{1,2}', s):
            raise ValueError('Tutar biçimi: 1234,56 veya 1234.56')
        s = s.replace('.', '').replace(',', '.')
    if not re.fullmatch(r'\d+(?:\.\d{1,2})?', s):
        raise ValueError('Pozitif tutar kullanın: 1234,56 veya 1234.56')
    try:
        return Decimal(s).quantize(CENT, rounding=ROUND_HALF_UP)
    except InvalidOperation:
        raise ValueError('Geçersiz tutar.') from None

def date_value(s):
    try:
        return datetime.strptime(s.strip(), '%d.%m.%Y')
    except ValueError:
        raise ValueError('Tarih GG.AA.YYYY biçiminde olmalı.') from None

def normalized(s):
    s = s.upper().replace('$', 'S')
    s = re.sub(r'\b(\d{2})[ \t]*([./-])[ \t]*(\d{2})[ \t]*[./-][ \t]*(\d{4})\b', r'\1\2\3\2\4', s)
    s = re.sub(r'(?m)^(\s*(?:TARIH\s*[:.]?\s*)?)(\d{2})[ ./-]+(\d{2})[ ./-]+(20\d{2})\b',r'\1\2.\3.\4',s)
    s = re.sub(r'\bARATOPLAM\b', 'ARA TOPLAM', s)
    # Yalnız para değerinin iki kuruş hanesinde OCR harflerini düzelt.
    s = re.sub(r'(\d+[.,])([0-9OC]{2})(?![A-Z0-9])', lambda m:m[1]+m[2].replace('O','0').replace('C','0'), s)
    s = ''.join(c for c in unicodedata.normalize('NFKD', s) if not unicodedata.combining(c))
    s = re.sub(r'\bTCIPLAM\b','TOPLAM',s)
    s = re.sub(r'\bTCIPKDV\b','TOPKDV',s)
    s = re.sub(r'\bF1S(?=[ \t]+N[O0])','FIS',s)
    s = re.sub(r'\bFI[.]S(?=[ \t]+N[O0])','FIS',s)
    s = re.sub(r'\b(FIS|FATURA|BELGE)[ \t]+N0\b',r'\1 NO',s)
    return re.sub(r'%[ \t]*:', '%', s)

def seller_name(s):
    # Yalnız sözcük içindeki OCR dolar işaretini düzelt; para simgelerine dokunma.
    s = re.sub(r'(?<=[A-Za-zÇĞİÖŞÜçğıöşü])\$(?=[A-Za-zÇĞİÖŞÜçğıöşü])', 'Ş', s)
    return re.sub(r'\bOCAKBA[SŞȘ]I\b', 'OCAKBAŞI', s, flags=re.IGNORECASE)

def ocr_money(s):
    # OCR bazen 5.900,00 değerini 5.900.00 okur.
    if re.fullmatch(r'\d{1,3}(?:\.\d{3})+\.\d{2}', s):
        left, cents = s.rsplit('.', 1)
        s = left.replace('.', '') + '.' + cents
    return money(s)

def gross_line(rate, gross):
    rate, gross = money(rate), money(gross)
    if rate > 100 or gross <= 0:
        raise ValueError('KDV oranı 0–100 arasında, dahil tutar sıfırdan büyük olmalı.')
    base = (gross / (1 + rate / 100)).quantize(CENT, rounding=ROUND_HALF_UP)
    return dict(oran=str(rate), matrah=str(base), kdv=str(gross-base))

ADDRESS_WORDS = r'\b(?:MAH|MAHALLESI|CAD|CADDESI|SOK|SOKAK|BULV|BULVARI|NO|TEL|TELEFON|ADANA|SEYHAN)\b'
BUSINESS_WORDS = r'\b(?:NOSTALJI|PETSHOP|AKVARYUM|AYAKKABI|DUNYASI|DUNVASI|MARKET|MAGAZA|MAGAZASI|LTD|STI|SAN|TIC|GIDA|OCAKBASI|RESTORAN|KAFE|KEBAPCI|TOBACCO|ELEKTRIK|TESEKKUR|EDERIZ)\b'

def receipt_header(text):
    result = []
    for line in text.splitlines():
        u = normalized(line)
        if re.search(r'\b(?:MAH|NAH|MH|MAHALLESI|CAD|CADDESI|SOK|SOKAK|BUL|BLV|BULV|BULVARI|NO|TEL|TELEFON|MEVKII|KOY)\b',u) or re.search(r'\d{2}[./-]\d{2}[./-]\d{4}|VKN|TCKN|\bTC\b|VERGI|\bV\.?D\b|FIS[ \t]*NO',u):
            break
        if re.search(r'\b(?:ARSIV|ARGIV|FNTURA|FATURA|MGZ|MAGAZA KODU)\b',u):
            result=[];continue
        if re.search(r'GOOGLE PLAY|APP STORE|UYGULAMA|HOS GELDINIZ|TESEKKUR',u):continue
        if (len(u.split()) >= 2 or re.search(r'\b(?:LTD|STI|A\.S)\b',u)) and (not re.search(r'\d',u) or re.search(r'\bA\.?\s*S\.?\s*$',u)):
            result.append(line.strip())
    return result

def person_name(text):
    candidates = [line for line in receipt_header(text)
                  if not re.search(BUSINESS_WORDS,normalized(line))
                  and re.fullmatch(r"[A-ZÇĞİÖŞÜÀ-Ža-zçğıöşü .'-]+",line)
                  and 2 <= len(line.split()) <= 4]
    return seller_name(candidates[0]) if len(candidates)==1 else ''

def unlabelled_header_ids(text):
    """Tarih/işlem öncesi tek başına basılmış, sağlama basamakları geçerli TCKN."""
    result=[]
    for line in text.splitlines():
        u=normalized(line)
        if re.search(r'\d{2}[./-]\d{2}[./-]\d{4}|\b(?:MUSTERI|ALICI|SATIS|BELGE|FIS|TOPLAM|ODEMELER|TERMINAL)\b',u):break
        if re.fullmatch(r'[1-9]\d{10}',u.strip()):
            n=u.strip();a=list(map(int,n))
            if (7*sum(a[:9:2])-sum(a[1:8:2]))%10==a[9] and sum(a[:10])%10==a[10]:result.append(n)
    return result


def relevant_text(text):
    from z_reports import is_z_report
    if is_z_report(text):return text
    kept = []
    header = set(receipt_header(text))
    for i,line in enumerate(text.splitlines()):
        u = normalized(line)
        if line.startswith(('[OCR KONTROL:', '[OCR BILGI:')):
            kept.append(line);continue
        # Tarih ve saat aynı satırda olabilir; yalnız saat parçasını çıkar.
        line = re.sub(r'\bSAAT\s*:\s*\d{1,2}[:.]\d{2}(?:[:.]\d{2})?', '', line, flags=re.IGNORECASE).strip()
        u = normalized(line)
        if re.search(r'\b(?:MUSTERI|MOSTERI|ALICI|NIHAI|TUKETICI)\b|11111111111',u):
            continue
        if re.search(r'\bV\.?\s*D\.?M?\.?\b|VERGI\s+DAIRESI',u):
            kept.append(line)
            following=text.splitlines()[i+1:i+2]
            if following and re.fullmatch(r'\s*\d{10,11}\s*',following[0]):
                kept.append(following[0])
            continue
        if re.search(r'\b(?:PARA)?PUAN\w*|IS\s*YERI\s*NO|TERMINAL|REF\s*NO|ONAY\s*KOD|KREDI\s*KART|CREDIT|\bAID\s*:|EKU\s*NO|Z\s*NO',u):
            continue
        if line.strip() in header or re.search(r'\d{2}[./-]\d{2}[./-]\d{4}|FIS\s*NO|FATURA\s*NO|BELGE\s*NO|VKN|TCKN|\bVN[ \t]*:|\bT\.?[ \t]*C\.?[ \t]*:|VERGI|\bV\.?[ \t]*D\.?[ \t]*[:.\-]?[ \t]*\d|TOPLAM|\bTOP\b|ODENECEK|KDV|%\s*\d|^[A-Z ]+/[ \t]*\d{10}\b',u):
            kept.append(line)
    return '\n'.join(kept)

def draft(text='', kind=None, use_vendor_hints=True):
    """Yalnızca etiketli adaylar çıkarılır; tüm sonuçlar insan onayı gerektirir."""
    from z_reports import is_z_report
    kind=kind or ('gelir' if is_z_report(text) else 'gider')
    from invoices import is_invoice, parse as parse_invoice
    if is_invoice(text):return parse_invoice(text,kind)
    from tax_directory import seller_ids
    annotations=[line for line in text.splitlines() if line.startswith(('[OCR KONTROL:', '[OCR BILGI:'))]
    text='\n'.join(line for line in text.splitlines() if not line.startswith(('[OCR KONTROL:', '[OCR BILGI:')))
    directory_ids=seller_ids(text)+unlabelled_header_ids(text)
    text = relevant_text(text)
    for number in set(directory_ids):
        if number not in text:text+='\nVKN: '+number
    d = dict(tarih='', evrak_tarihi='', aciklama='', vkn='', evrak_no='',
             kod='ML' if kind=='gelir' else 'MS', toplam='', kalemler=[], ocr=text, kind=kind)
    lines = [x.strip() for x in text.splitlines() if x.strip()]
    upper = normalized(text)
    date_header=re.split(r'\b(?:TOPLAM|TOPKDV|KREDI|BANKA)\b',upper)[0]
    # Fiş numarasından sonraki POS tarihlerini belge tarihi adayına katma.
    header_lines=[]
    for line in date_header.splitlines():
        header_lines.append(line)
        if re.search(r'\bFIS[ \t]*NO',line) and re.search(r'\d{2}[./-]\d{2}[./-]\d{4}','\n'.join(header_lines)):
            break
    date_header='\n'.join(header_lines)
    dates = re.findall(r'\b(\d{2})[./-](\d{2})[./-](\d{4})\b', date_header)
    unique = list(dict.fromkeys('.'.join(x) for x in dates))
    if len(unique) == 1:
        try:
            date_value(unique[0])
            d['tarih'] = d['evrak_tarihi'] = unique[0]
        except ValueError:
            pass
    ids = re.findall(r'(?:VKN|\bVN|VERGI[ \t]*(?:KIMLIK[ \t]*)?NO|TCKN|\bT\.?[ \t]*C\.?(?:[ \t]*(?:KIMLIK[ \t]*)?NO)?|\bV\.?[ \t]*D\.?M?\.?|VERGI[ \t]+DAIRESI)[ \t]*[:.\-]?[ \t]*(?:\n[ \t]*)?(\d{10,11})(?!\d)', upper)
    ids=[i for i in ids+directory_ids if i!='11111111111']
    # Vergi dairesi adı / VKN biçimi yalnız belge tarihinden önceki başlıkta.
    before_document=re.split(r'\d{2}[./-]\d{2}[./-]\d{4}|\b(?:FIS|BELGE)[ \t]+NO',upper)[0]
    ids += re.findall(r'\b(?:MM|MAL\s*MUDURLUGU)\s*[:.]\s*(\d{10})(?!\d)',before_document)
    ids += re.findall(r'(?m)^[A-Z ]+/[ \t]*(\d{10})\b',before_document)
    if len(set(ids)) == 1:
        d['vkn'] = ids[0]
    elif len(set(ids))>1:
        annotations.append('[OCR KONTROL: Birden fazla satıcı kimliği okundu; görüntüde birden fazla fiş olabilir. Belgeleri ayrı kontrol edin.]')
    nums = []
    for label in ('FIS','FATURA','BELGE'):
        found = re.findall(r'\b'+label+r'[ \t]*(?:NO|NUMARASI)[ \t]*[:.\-]?[ \t]*([A-Z0-9/.:\-]+)',upper)
        if label=='FIS':found=[re.sub(r'[.:]','',n) if re.fullmatch(r'\d+(?:[.:]\d+)*[.]?',n) else n for n in found]
        found = [n for n in found if (n.isdigit() if label=='FIS' else any(c.isdigit() for c in n))]
        if found:
            nums=found
            break
    if len(set(nums)) == 1:
        d['evrak_no'] = str(int(nums[0])) if nums[0].isdigit() else nums[0]
        if nums[0].isdigit() and len(nums[0])>16:
            d['original_evrak_no']=nums[0]
            d['evrak_no']=nums[0][4:]
    # İlk satır yalnızca düzenlenebilir açıklama önerisidir.
    header = receipt_header(text)
    # Ticaret unvanı birden çok satıra basılır; yalnız LTD/ŞTİ satırını seçme.
    legal=[i for i,line in enumerate(header) if re.search(r'\bA\s*\.\s*S\.?\s*$|\bLTD\b|\bLIMITED\b|\bANONIM\b',normalized(line))]
    company=' '.join(header[:legal[-1]+1]) if legal else ' '.join(header)
    d['aciklama'] = person_name(text) if len(d['vkn'])==11 else seller_name(company or (header[0] if header else ''))
    # Kullanıcının belge üzerinden doğruladığı unvan; başka firmalara harf tahmini yapma.
    verified = json.loads((Path(__file__).parent/'verified_vendors.json').read_text(encoding='utf-8'))
    d['ocr_identity']={'vkn':d['vkn'],'aciklama':d['aciklama']}
    if use_vendor_hints and d['vkn'] in verified:
        d['aciklama'] = verified[d['vkn']]
    totals = [];payables=[];intermediate=[]
    for line in lines:
        u = normalized(line)
        if re.match(r'^(?:(?:GENEL|ARA)\s+TOPLAM|TOPLAM|TOP|ODENECEK\s+(?:KDV\s+DAHIL\s+)?TUTAR)\s*[:* ]', u) and ('KDV' not in u or u.startswith('ODENECEK')):
            vals = re.findall(r'\d+(?:[.,]\d{3})*[.,]\d{2}(?!\d)', u)
            if vals:
                try:
                    if u.startswith('ARA '):
                        intermediate.append(str(ocr_money(vals[-1])));continue
                    totals.append(str(ocr_money(vals[-1])))
                    if u.startswith('ODENECEK'):payables.append(totals[-1])
                except ValueError:
                    pass
    if not totals:totals=intermediate
    if len(set(payables))==1:totals=payables
    if len(set(totals)) == 1:
        d['toplam'] = totals[0]
    # Yalnız tek açık oran ve etiketli toplam varsa dahil tutardan aday hesaplanır.
    rates = set(re.findall(r'%\s*(\d{1,2})(?!\d)(?:[.,]0{1,2}|[.])?\s+[*× ]*\d+(?:[.,]\d{3})*[.,]\d{2}(?!\d)', upper))
    if not rates:
        rates=set(re.findall(r'(?m)^(?:KDV\s*)?%\s*(\d{1,2})\s*$',upper))
    if d['toplam'] and len(rates) == 1:
        try:
            d['kalemler'] = [gross_line(next(iter(rates)),d['toplam'])]
        except ValueError:
            pass
    elif d['toplam'] and len(rates)>1:
        # Yalnız fişteki KDV özet tablosunun oran/dahil-tutar satırları kullanılır.
        summary = False
        last_amount = False
        groups = []
        for line in lines:
            u = normalized(line)
            if 'DAHIL' in u and 'KDV' in u:
                summary = True
                last_amount = 'MATRAH' in u
                continue
            if summary:
                m=re.match(r'^[ \t]*%[ \t]*(\d{1,2})(?!\d)(?:[.,]\d{2}|[.])?[ \t]+[* ]*([\d.,]+)',u)
                if m:
                    try:
                        amounts=re.findall(r'\d+(?:[.,]\d{3})*[.,]\d{2}(?!\d)',u[m.end(1):])
                        value=amounts[-1] if last_amount and amounts else m[2]
                        groups.append(gross_line(m[1],str(ocr_money(value))))
                    except ValueError:
                        pass
                elif groups:
                    break
        if groups and sum(money(g['matrah'])+money(g['kdv']) for g in groups)==money(d['toplam']):
            d['kalemler']=groups
    # Özet yoksa yalnız açık oran/tutar içeren ürün satırlarını topla.
    # İndirim veya okunmayan satır varsa toplam tutmayacağı için tahmin üretme.
    if not d['kalemler'] and d['toplam'] and len(rates)>1:
        grouped={}
        for line in lines:
            u=normalized(line)
            # Yalnız ürün açıklaması bulunan satırlar; özet satırlarını ikinci kez toplama.
            if not re.search(r'[A-Z]',u.split('%')[0]) or re.search(r'KDV|TOPLAM|DAHIL',u):continue
            m=re.search(r'%\s*(\d{1,2})(?!\d)(?:[.,]0{1,2}|[.])?\s*[*× ]+(-?[\d.,]+)',u)
            if m:
                try:
                    rate=str(int(m[1])); grouped[rate]=grouped.get(rate,Decimal(0))+(Decimal(-1) if m[2].startswith('-') else Decimal(1))*ocr_money(m[2].lstrip('-'))
                except ValueError:
                    pass
        if grouped and sum(grouped.values())==money(d['toplam']):
            d['kalemler']=[gross_line(r,str(g)) for r,g in grouped.items()]
    if not d['kalemler'] and d['toplam'] and not rates:
        # Basılı KDV ve brüt tutar tek bir oranı doğrularsa inceleme önerisi üret.
        vat_values=[]
        for line in lines:
            m=re.fullmatch(r'(?:TOPKDV|TOPLAM KDV|KDV)\s*[:*+. ]+([\d.,]+)',normalized(line))
            if m:
                try:vat_values.append(ocr_money(m[1]))
                except ValueError:pass
        if len(set(vat_values))==1:
            candidates=[gross_line(r,d['toplam']) for r in (1,10,20)]
            candidates=[g for g in candidates if abs(money(g['kdv'])-vat_values[0])<=CENT]
            if len(candidates)==1:
                d['kalemler']=candidates
                d['review_note']='KDV oranı metinde bulunamadı; basılı KDV ve toplamdan hesaplandı. Fotoğraftan doğrulayın.'
                d['vat_inferred']=True
    from z_reports import is_z_report, parse
    if kind=='gelir' and is_z_report(text):
        z=parse(text);warning=z.pop('warning',None);d.update(z)
        d['vkn']=''
        d.pop('vat_inferred',None);d.pop('review_note',None)
        if warning:annotations.append('[OCR KONTROL: '+warning+']')
    zero=sum(money(g['matrah'])+money(g['kdv']) for g in d['kalemler'] if money(g['oran'])==0)
    if zero and kind!='gelir':
        d['excluded_zero_total']=str(zero)
        d['kalemler']=[g for g in d['kalemler'] if money(g['oran'])!=0]
    if kind=='gelir':d['vkn']=''
    if annotations:d['ocr']+='\n'+'\n'.join(annotations)
    from fis_cikarim import tamamla
    d = tamamla(d, text, kind)
    return d

def validate(d, *, check_person_name=True):
    if d.get('mixed_documents'):raise ValueError('Birden fazla Z raporu birleşmiş. Her raporu ayrı kırpıp yükleyin.')
    if d.get('kind','gider') not in ('gelir','gider'):
        raise ValueError('Gelir / gider türü geçersiz.')
    if d.get('kind')=='gelir' and d.get('kod')=='MS':
        raise ValueError('MS masraf kodu gelir kaydında kullanılamaz.')
    optional_id=d.get('kind')=='gelir' and not d.get('is_invoice')
    if d.get('kind')=='gelir' and not d.get('is_invoice') and d.get('vkn'):
        raise ValueError('Gelirde VKN / TCKN yalnız faturada kullanılır; Z raporu kimliği boş olmalı.')
    for k, label in [('aciklama','Açıklama'), ('evrak_no','Fiş / evrak no')]+([] if optional_id else [('vkn','VKN / TCKN')]):
        if not d.get(k, '').strip():
            raise ValueError(label + ' eksik.')
    for k in ('tarih', 'evrak_tarihi'):
        date_value(d.get(k, ''))
    if (d.get('vkn') or not optional_id) and (not re.fullmatch(r'\d{10}|\d{11}', d.get('vkn','')) or d.get('vkn')=='11111111111'):
        raise ValueError('VKN 10, TCKN 11 rakam olmalı. Bu kontrol kimliği doğrulamaz.')
    if check_person_name and d.get('kind')!='gelir' and not d.get('custom_description') and len(d.get('vkn',''))==11 and (len(d['aciklama'].split())<2 or re.search(BUSINESS_WORDS,normalized(d['aciklama']))):
        raise ValueError('TCKN bulunan fişte açıklama alanına işletme tabelası değil kişinin adını ve soyadını yazın.')
    if d.get('kod') not in (('ML','HZ') if d.get('kind')=='gelir' else ('MS','ML')):
        raise ValueError('Gelirde ML veya HZ; giderde MS veya ML seçin.')
    if not d.get('kalemler'):
        raise ValueError('En az bir KDV satırı girin.')
    evidence=draft(d.get('ocr',''),d.get('kind','gider'))
    observed = {money(r['oran']) for r in evidence['kalemler'] if money(r['oran'])!=0 or d.get('kind')=='gelir'}
    entered = {money(row.get('oran','')) for row in d['kalemler']}
    if d.get('vat_partition') is not None:
        partition=money(d['vat_partition'])
        if entered!={partition}:
            raise ValueError('Bu kayıt yalnız kendi KDV oranını içerebilir.')
        observed=observed & {partition}
    if observed and observed != entered and not d.get('vat_manually_reviewed'):
        shown = ', '.join('%'+format(x,'g') for x in sorted(observed))
        raise ValueError(f'Fişte okunan KDV oranı {shown}; tablodaki oran farklı. Fotoğrafı kontrol edip KDV satırını düzeltin.')
    total = Decimal(0)
    for row in d['kalemler']:
        rate, base, vat = (money(row.get(k, '')) for k in ('oran','matrah','kdv'))
        if rate > 100 or base <= 0:
            raise ValueError('KDV oranı 0–100 arasında, matrah sıfırdan büyük olmalı.')
        expected = (base * rate / 100).quantize(CENT, rounding=ROUND_HALF_UP)
        if abs(expected-vat) > Decimal('0.02'):
            raise ValueError(f'Matrah / KDV uyuşmuyor: {base} × %{rate} = {expected}, girilen {vat}.')
        total += base + vat
    if abs(total+money(d.get('excluded_zero_total','0'))-money(d.get('toplam',''))) > Decimal('0.02'):
        raise ValueError(f'Fiş toplamı uyuşmuyor. KDV satırları toplamı: {total}')
    printed=re.findall(r'(?m)^(?:TOPKDV|TOPLAM KDV|KDV)\s*[:*×+. ]*(\d+(?:[.,]\d{3})*[.,]\d{2})(?!\d)',normalized(d.get('ocr','')))
    if len(set(printed))==1 and d.get('vat_partition') is None and not d.get('vat_manually_reviewed'):
        if abs(sum(money(r['kdv']) for r in d['kalemler'])-ocr_money(printed[0]))>Decimal('0.03'):
            raise ValueError('Fişte basılı toplam KDV ile hesaplanan KDV uyuşmuyor. Tutar ve oranı fotoğraftan kontrol edin.')

def validate_export(d, rid):
    """Aylık çıktı önceden onaylı kayıtlardır; yeni ünvan sezgisini geriye uygulama."""
    try:validate(d,check_person_name=False)
    except ValueError as error:
        label='Gelir' if d.get('kind')=='gelir' else 'Gider'
        raise ValueError(f"{label} kaydı #{rid} — {d.get('evrak_tarihi','')} / No: {d.get('evrak_no','')} / {d.get('aciklama','')}: {error}") from error

def identity(d):
    key='|'.join([d.get('vkn',''), date_value(d['evrak_tarihi']).strftime('%Y-%m-%d'), d['evrak_no'].strip().upper()])
    if d.get('kind')=='gelir':key='GELIR|'+key
    return key+'|KDV:'+str(money(d['vat_partition'])) if d.get('vat_partition') is not None else key

def vat_records(d):
    """Tek belgeden her pozitif KDV oranı için bağımsız kayıt üret."""
    groups={}
    for row in d.get('kalemler',[]):
        rate=money(row['oran'])
        if rate or d.get('kind')=='gelir':
            groups[rate]=groups.get(rate,Decimal(0))+money(row['matrah'])+money(row['kdv'])
    if len(groups)<2:
        return [d]
    # Tutarı tutmayan belgeyi bölerek hatayı gizleme.
    if abs(sum(groups.values())+money(d.get('excluded_zero_total','0'))-money(d.get('toplam','0')))>CENT:
        return [d]
    results=[]
    for rate,total in groups.items():
        part=dict(d)
        part.update(vat_partition=str(rate),source_receipt_total=d.get('source_receipt_total',d['toplam']),
                    toplam=str(total),kalemler=[dict(oran=str(rate),matrah=str(sum(money(r['matrah']) for r in d['kalemler'] if money(r['oran'])==rate)),kdv=str(sum(money(r['kdv']) for r in d['kalemler'] if money(r['oran'])==rate)))])
        part.pop('excluded_zero_total',None)
        results.append(part)
    return results

def posting_period(d,period):
    d=dict(d)
    if period and d.get('evrak_tarihi'):
        start=datetime.strptime(period,'%m-%Y')
        try:
            document=date_value(d['evrak_tarihi'])
        except ValueError:
            return d
        d['tarih']=(start if document<start and not document_month(d) else document).strftime('%d.%m.%Y')
    return d

def document_month(d):
    """Gelirler ve faturalar kendi belge tarihleriyle kaydedilir."""
    return bool(d.get('is_z_report') or d.get('is_invoice') or d.get('kind')=='gelir')

def actual_period(d, fallback):
    if document_month(d) and d.get('evrak_tarihi'):
        try:return date_value(d['evrak_tarihi']).strftime('%m-%Y')
        except ValueError:pass
    return fallback

class Store:
    def __init__(self, path):
        self.db = sqlite3.connect(path)
        self.db.execute('PRAGMA foreign_keys=ON')
        self.db.execute('PRAGMA busy_timeout=10000')
        self.db.executescript('''
            CREATE TABLE IF NOT EXISTS accounting_profiles(client INTEGER PRIMARY KEY REFERENCES clients(id), data TEXT NOT NULL);
            CREATE TABLE IF NOT EXISTS monthly_exports(client INTEGER, period TEXT, path TEXT NOT NULL, PRIMARY KEY(client,period));
            CREATE TABLE IF NOT EXISTS clients(id INTEGER PRIMARY KEY, name TEXT UNIQUE NOT NULL);
            CREATE TABLE IF NOT EXISTS vendor_aliases(alias TEXT, vkn TEXT, PRIMARY KEY(alias,vkn));
            CREATE TABLE IF NOT EXISTS deleted_vendors(vkn TEXT PRIMARY KEY);
            CREATE TABLE IF NOT EXISTS vendors(vkn TEXT PRIMARY KEY, name TEXT NOT NULL);
            CREATE TABLE IF NOT EXISTS receipts(
              id INTEGER PRIMARY KEY, client INTEGER REFERENCES clients(id),
              image TEXT, hash TEXT, data TEXT NOT NULL, approved INTEGER DEFAULT 0,
              identity TEXT, exported TEXT,
              UNIQUE(client,hash));
            CREATE UNIQUE INDEX IF NOT EXISTS unique_document ON receipts(client,identity)
            WHERE identity IS NOT NULL;
        ''')
        if 'period' not in [r[1] for r in self.db.execute('PRAGMA table_info(receipts)')]:
            self.db.execute('ALTER TABLE receipts ADD COLUMN period TEXT')
            self.db.commit()
        if 'archived' not in [r[1] for r in self.db.execute('PRAGMA table_info(clients)')]:
            self.db.execute('ALTER TABLE clients ADD COLUMN archived INTEGER NOT NULL DEFAULT 0')
            self.db.commit()
        self.db.execute('CREATE INDEX IF NOT EXISTS receipts_month ON receipts(client,period,approved)')
        self.db.execute('CREATE INDEX IF NOT EXISTS aliases_vkn ON vendor_aliases(vkn)')
        self.db.execute('CREATE TABLE IF NOT EXISTS app_migrations(name TEXT PRIMARY KEY)')
        with self.db:
            seed=not self.db.execute("SELECT 1 FROM app_migrations WHERE name='vendor_history_v1'").fetchone()
            historical=self.db.execute('SELECT data FROM receipts WHERE approved=1 ORDER BY id DESC').fetchall() if seed else []
            for (data,) in historical:
                d=json.loads(data)
                if d.get('kind')!='gelir' and not d.get('custom_description') and d.get('vkn') and d.get('aciklama') and not self.db.execute('SELECT 1 FROM deleted_vendors WHERE vkn=?',(d['vkn'],)).fetchone():
                    self.db.execute('INSERT OR IGNORE INTO vendors VALUES (?,?)',(d['vkn'],d['aciklama']))
                    self.learn_vendor(d)
            self.db.execute("INSERT OR IGNORE INTO app_migrations VALUES ('vendor_history_v1')")

    def delete_vendor(self,vkn):
        with self.db:
            self.db.execute('INSERT OR IGNORE INTO deleted_vendors VALUES (?)',(vkn,))
            self.db.execute('DELETE FROM vendors WHERE vkn=?',(vkn,))
            self.db.execute('DELETE FROM vendor_aliases WHERE vkn=?',(vkn,))

    def learn_vendor(self,d):
        if d.get('kind')=='gelir' or d.get('custom_description') or not d.get('vkn'):return
        aliases=receipt_header(d.get('ocr',''))+[d['aciklama']]
        for name in aliases:
            key=re.sub(r'[^A-Z0-9]','',normalized(name))
            if len(key)>=5 and len(name.split())>=2 and key not in {'TESEKKUREDERIZ','AFIYETOLSUN','HOSGELDINIZ','MUSTERINUSHASI','BUBELGEYISAKLAYINIZ','KREDIKARTI','EARSIVFATURA','BILGIFISI'}:
                self.db.execute('INSERT OR IGNORE INTO vendor_aliases VALUES (?,?)',(key,d['vkn']))

    def recall_vendor(self,d):
        if d.get('kind')=='gelir' or d.get('is_invoice') or d.get('custom_description') or d.get('vkn_omitted'):return d
        row=self.db.execute('SELECT name FROM vendors WHERE vkn=?',(d.get('vkn',''),)).fetchone()
        if row:
            d['aciklama']=row[0];d['vendor_source']='Ünvan: kayıtlı cariden alındı (VKN eşleşmesi). VKN: '+('önceki cari eşleşmesinden' if d.get('identity_from_memory') else 'belge alanından')
        elif not d.get('vkn') or d.get('vkn')=='11111111111':
            def key(name):return re.sub(r'[^A-Z0-9]','',normalized(name))
            names={key(d.get('aciklama',''))}
            names.update(key(n) for n in receipt_header(d.get('ocr','')))
            matches={v:n for v,n in self.vendor_rows() if key(n) and key(n) in names}
            for alias in names:
                for v,n in self.db.execute('SELECT v.vkn,v.name FROM vendor_aliases a JOIN vendors v ON v.vkn=a.vkn WHERE a.alias=?',(alias,)):
                    matches[v]=n
            if len(matches)==1:
                d['vkn'],d['aciklama']=next(iter(matches.items()));d['identity_from_memory']=True;d['vendor_source']='VKN ve ünvan: onaylı ünvan eşleşmesinden tamamlandı'
            elif len(matches)>1:d['vendor_source']='Birden fazla cari eşleşti; kimlik seçimi gerekli'
        return d

    def suggest_vendors(self,d):
        from difflib import SequenceMatcher
        def key(x):return re.sub(r'[^A-Z0-9]','',normalized(x))
        names={key(d.get('aciklama',''))}
        names.update(key(n) for n in receipt_header(d.get('ocr','')))
        names={n for n in names if len(n)>=5}
        candidates=[];alias_map={}
        for alias,vkn in self.db.execute('SELECT alias,vkn FROM vendor_aliases'):alias_map.setdefault(vkn,[]).append(alias)
        for v,n in self.vendor_rows():
            aliases=[key(n)]+alias_map.get(v,[])
            score=max((SequenceMatcher(None,a,b).ratio() for a in names for b in aliases),default=0)
            if score>=.75:candidates.append((score,v,n))
        return [(v,n) for _,v,n in sorted(candidates,reverse=True)[:5]]

    def clients(self,archived=False):
        return sorted(self.db.execute('SELECT id,name FROM clients WHERE archived=?',(int(archived),)).fetchall(),key=lambda r:turkish_sort(r[1]))

    def rename_client(self,cid,name):
        if not name.strip(): raise ValueError('Mükellef adı boş olamaz.')
        with self.db:self.db.execute('UPDATE clients SET name=? WHERE id=?',(name.strip(),cid))

    def archive_client(self,cid,archived=True):
        with self.db:self.db.execute('UPDATE clients SET archived=? WHERE id=?',(int(archived),cid))

    def vendor_rows(self,query=''):
        words=normalized(query).split()
        rows=self.db.execute('SELECT vkn,name FROM vendors WHERE vkn<>?',('11111111111',)).fetchall()
        return sorted([r for r in rows if all(w in normalized(r[0]+' '+r[1]) for w in words)],key=lambda r:turkish_sort(r[1]))

    def update_vendor(self,old_vkn,vkn,name):
        if not re.fullmatch(r'\d{10,11}',vkn) or vkn=='11111111111':raise ValueError('Geçerli uzunlukta VKN/TCKN girin; 11111111111 kullanılamaz.')
        if not name.strip():raise ValueError('Ünvan boş olamaz.')
        with self.db:
            if old_vkn!=vkn:
                self.db.execute('INSERT INTO vendors VALUES (?,?)',(vkn,name.strip()))
                self.db.execute('DELETE FROM vendors WHERE vkn=?',(old_vkn,))
                self.db.execute('UPDATE OR IGNORE vendor_aliases SET vkn=? WHERE vkn=?',(vkn,old_vkn))
                self.db.execute('DELETE FROM vendor_aliases WHERE vkn=?',(old_vkn,))
            else:self.db.execute('INSERT OR REPLACE INTO vendors VALUES (?,?)',(vkn,name.strip()))

    def add_vendor(self,vkn,name):
        if self.db.execute('SELECT 1 FROM vendors WHERE vkn=?',(vkn,)).fetchone():
            raise ValueError('Bu VKN/TCKN zaten kayıtlı. Düzenle seçeneğini kullanın.')
        self.update_vendor(vkn,vkn,name)

    def delete_receipts(self,client,ids):
        ids=list(dict.fromkeys(int(i) for i in ids))
        periods=set()
        with self.db:
            self.db.execute('CREATE TABLE IF NOT EXISTS deleted_receipts(id INTEGER PRIMARY KEY, deleted_at TEXT, payload TEXT)')
            for rid in ids:
                cursor=self.db.execute('SELECT * FROM receipts WHERE id=? AND client=?',(rid,client))
                row=cursor.fetchone()
                if row is None:raise ValueError('Seçilen kayıt bu mükellefe ait değil veya silinmiş.')
                payload=dict(zip([c[0] for c in cursor.description],row))
                try:periods.add(self.receipt_period(rid))
                except ValueError:pass
                self.db.execute('INSERT INTO deleted_receipts(deleted_at,payload) VALUES (?,?)',(datetime.now().isoformat(),json.dumps(payload,ensure_ascii=False)))
                self.db.execute('DELETE FROM receipts WHERE id=?',(rid,))
        return periods

    def client(self, name):
        name = name.strip()
        if not name:
            raise ValueError('Mükellef adı boş olamaz.')
        with self.db:
            self.db.execute('INSERT OR IGNORE INTO clients(name) VALUES (?)', (name,))
            self.db.execute('UPDATE clients SET archived=0 WHERE name=?',(name,))
        return self.db.execute('SELECT id FROM clients WHERE name=?',(name,)).fetchone()[0]

    def add(self, client, image, data, period=None):
        from accounting import apply_defaults
        data=self.recall_vendor(apply_defaults(self,client,dict(data)))
        period=actual_period(data,period)
        if period:
            datetime.strptime(period, '%m-%Y')
            data=posting_period(data,period)
        digest = hashlib.sha256(Path(image).read_bytes()).hexdigest() if image else None
        with self.db:
            rid=self.db.execute('INSERT INTO receipts(client,image,hash,data,period) VALUES (?,?,?,?,?)',
                (client,str(image),digest,json.dumps(data,ensure_ascii=False),period)).lastrowid
            self._split_saved(rid,data,False)
            return rid

    def _split_saved(self,rid,d,approve):
        parts=vat_records(d)
        if len(parts)<2: return
        client,image,period=self.db.execute('SELECT client,image,period FROM receipts WHERE id=?',(rid,)).fetchone()
        for i,part in enumerate(parts):
            key=identity(part) if approve else None
            if i==0:
                self.db.execute('UPDATE receipts SET data=?,identity=? WHERE id=?',(json.dumps(part,ensure_ascii=False),key,rid))
            else:
                self.db.execute('INSERT INTO receipts(client,image,data,period,approved,identity) VALUES (?,?,?,?,?,?)',
                    (client,image,json.dumps(part,ensure_ascii=False),period,int(approve),key))

    def receipt_period(self, rid):
        period, data = self.db.execute('SELECT period,data FROM receipts WHERE id=?',(rid,)).fetchone()
        return period or date_value(json.loads(data)['tarih']).strftime('%m-%Y')

    def monthly_records(self, client, period):
        datetime.strptime(period,'%m-%Y')
        return [(i,posting_period(d,period)) for i,d,a,e in reversed(self.rows(client,period=period,approved_only=True))]

    def month_path(self,client,period,default=None):
        row=self.db.execute('SELECT path FROM monthly_exports WHERE client=? AND period=?',(client,period)).fetchone()
        return Path(row[0]) if row else (Path(default) if default else None)

    def sync_month(self, client, period, path=None):
        if path is None:
            row = self.db.execute('SELECT path FROM monthly_exports WHERE client=? AND period=?',(client,period)).fetchone()
            if not row:
                return None
            path = row[0]
        path = Path(path)
        path.parent.mkdir(parents=True,exist_ok=True)
        records = self.monthly_records(client,period)
        from accounting import profile, export
        settings=profile(self,client)
        if settings['mode']=='genel':export(path,records,settings)
        else:export_monthly(path,records,settings)
        with self.db:
            self.db.execute('INSERT OR REPLACE INTO monthly_exports VALUES (?,?,?)',(client,period,str(path)))
        return path

    def rows(self, client, kind=None, period=None, approved_only=False):
        sql='SELECT id,data,approved,exported,period FROM receipts WHERE client=?'
        args=[client]
        if period:
            datetime.strptime(period,'%m-%Y')
            sql+=' AND (period=? OR period IS NULL OR period=?)';args.extend([period,''])
        if approved_only:sql+=' AND approved=1'
        rows=[]
        for i,payload,a,e,saved in self.db.execute(sql+' ORDER BY id DESC',args):
            d=posting_period(json.loads(payload),saved)
            if period and not saved:
                try:
                    if date_value(d.get('tarih','')).strftime('%m-%Y')!=period:continue
                except ValueError:continue
            if kind is None or d.get('kind','gider')==kind:rows.append((i,d,a,e))
        return rows

    def get(self, rid):
        row = self.db.execute('SELECT image,data,approved,exported,period FROM receipts WHERE id=?',(rid,)).fetchone()
        return row[0], posting_period(json.loads(row[1]),row[4]), row[2], row[3]

    def save(self, rid, d, approve=False):
        if self.get(rid)[3]:
            raise ValueError('Excel’e çıkarılan kayıt kilitli. Düzeltmeyi Zirve kaydıyla birlikte değerlendirin.')
        period=actual_period(d,self.receipt_period(rid))
        d=posting_period(d,period) if d.get('evrak_tarihi') else dict(d)
        old=self.get(rid)[1]
        if old.get('vkn') and not d.get('vkn'):d['vkn_omitted']=True
        elif d.get('vkn'):d.pop('vkn_omitted',None)
        from accounting import profile
        cid=self.db.execute('SELECT client FROM receipts WHERE id=?',(rid,)).fetchone()[0]
        d['accounting_mode']=profile(self,cid)['mode']
        key = None
        if approve:
            from review import warnings
            if warnings(d) and not d.get('review_ack'):raise ValueError('Okuma uyarılarını kontrol edip doğrulama kutusunu işaretleyin.')
            validate(d)
            from accounting import profile,resolve
            cid=self.db.execute('SELECT client FROM receipts WHERE id=?',(rid,)).fetchone()[0]
            settings=profile(self,cid)
            if settings['mode']=='genel':
                for line in d['kalemler']:resolve(d,settings,line['oran'])
            # Tek satırlı ve KDV'ye bölünmüş kopyalar aynı belgeyi ikinci kez aktaramaz.
            document_key=identity(dict(d,vat_partition=None))
            rates={money(r['oran']) for r in d['kalemler']}
            for other_id,payload in self.db.execute('SELECT id,data FROM receipts WHERE client=? AND approved=1 AND id<>?',(cid,rid)):
                other=json.loads(payload)
                if identity(dict(other,vat_partition=None))==document_key and rates & {money(r['oran']) for r in other.get('kalemler',[])}:
                    raise sqlite3.IntegrityError('Bu belge ve KDV oranı zaten onaylı.')
            key = identity(d)
        with self.db:
            self.db.execute('UPDATE receipts SET data=?,approved=?,identity=?,period=? WHERE id=?',
                (json.dumps(d,ensure_ascii=False),int(approve),key,period,rid))
            if approve and d.get('vkn') and d.get('kind')!='gelir' and not d.get('custom_description'):
                self.db.execute('INSERT OR REPLACE INTO vendors VALUES (?,?)',(d['vkn'],d['aciklama']))
                self.learn_vendor(d)
            self._split_saved(rid,d,approve)

    def pending(self, client, period):
        datetime.strptime(period, '%m-%Y')
        return [(i,d) for i,d,a,e in self.rows(client) if a and not e
                and self.receipt_period(i)==period]

    def mark_exported(self, ids, path):
        with self.db:
            self.db.executemany('UPDATE receipts SET exported=? WHERE id=?',[(str(path),i) for i in ids])

HEADERS = ['Kod','Evrak Tarihi','Evrak No','Vergi / TC No','Gider Açıklama','KDV %','Alınan Mal/Masraf']
INCOME_HEADERS = ['Kod','Evrak Tarihi','Evrak No','Gelir Açıklama','KDV %','Satılan Mal/Hizmet']

def export_xlsx(path, records, headers=True, allow_empty=False, kind='gider', workbook=None, sheet_name='Isletme', save=True):
    """Mevcut uygulamanın openpyxl dışa aktarımını işletme defterine uyarlar."""
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter
    if not records and not allow_empty:
        raise ValueError('Bu dönemde onaylı, henüz çıkarılmamış fiş yok.')
    wb = workbook if workbook is not None else Workbook()
    ws = wb.create_sheet(sheet_name) if workbook is not None else wb.active
    ws.title = sheet_name
    if headers:
        ws.append(INCOME_HEADERS if kind=='gelir' else HEADERS)
    for rid, d in records:
        validate_export(d,rid)
        for line in d['kalemler']:
            if money(line['oran'])==0 and kind!='gelir':
                continue
            values=[d['kod'],date_value(d['evrak_tarihi'] if kind=='gelir' else d['tarih']),d['evrak_no']]
            if kind!='gelir':values.append(d.get('vkn',''))
            ws.append(values+[d['aciklama'],float(money(line['oran'])),float(money(line['matrah'])+money(line['kdv']))])
            r = ws.max_row
            for c in ((1,3,4) if kind=='gelir' else (1,3,4,5)):
                ws.cell(r,c).data_type = 's'
                ws.cell(r,c).number_format = '@'
            for c in (2,):
                ws.cell(r,c).number_format = 'dd.mm.yyyy'
            for c in ((5,6) if kind=='gelir' else (6,7)):
                ws.cell(r,c).number_format = '#,##0.00'
    if headers:
        for c in ws[1]:
            c.font = Font(bold=True,color='FFFFFF')
            c.fill = PatternFill('solid',fgColor='234B58')
    for i,w in enumerate([10,16,24,42,12,25] if kind=='gelir' else [10,16,24,20,42,12,25],1):
        ws.column_dimensions[get_column_letter(i)].width = w
    if headers:
        ws.freeze_panes = 'A2'
        ws.auto_filter.ref = ws.dimensions
    if not save:
        return wb
    from excel_save import save_workbook
    save_workbook(wb,path)
    return wb

def export_monthly(path,records,settings=None):
    wb = None
    specs = [('gider','Giderler',True),('gelir','Gelirler',True),
             ('gider','Gider_Aktarim',False),('gelir','Gelir_Aktarim',False)]
    for index,(kind,name,headers) in enumerate(specs):
        selected = [(i,dict(d,kod=settings.get('income_code','ML')) if kind=='gelir' and settings else d) for i,d in records if d.get('kind','gider')==kind]
        wb = export_xlsx(path,selected,headers=headers,allow_empty=True,kind=kind,
                         workbook=wb,sheet_name=name,save=index==len(specs)-1)

