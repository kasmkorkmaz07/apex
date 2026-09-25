import unittest,tempfile,json
from pathlib import Path
from copy import deepcopy
from openpyxl import load_workbook
from core import Store,validate,validate_export
from accounting import EMIR,save_profile
from test_core import sample

class ApprovedExportTests(unittest.TestCase):
 def test_legacy_name_does_not_block_delete_sync(self):
  for mode in ('genel','isletme'):
   with self.subTest(mode=mode),tempfile.TemporaryDirectory() as tmp:
    s=Store(Path(tmp)/'db')
    try:
     cid=s.client('T')
     if mode=='genel':save_profile(s,cid,deepcopy(EMIR))
     d=sample();d.update(vkn='16171222776',aciklama='KEBAPÇI GÖKHAN')
     rid=s.add(cid,None,d,'09-2026')
     with self.assertRaisesRegex(ValueError,'TCKN bulunan'):s.save(rid,s.get(rid)[1],True)
     # Eski sürümde onaylanmış kayıt; yeni ad sezgisi Excel yenilemeyi durdurmamalı.
     with s.db:s.db.execute('UPDATE receipts SET approved=1 WHERE id=?',(rid,))
     other=s.add(cid,None,sample(no='99'),'09-2026');s.save(other,s.get(other)[1],True)
     path=Path(tmp)/'out.xlsx';s.sync_month(cid,'09-2026',path)
     s.delete_receipts(cid,[other]);s.sync_month(cid,'09-2026')
     w=load_workbook(path)
     try:
      ws=w['Genel_Alis_Aktarim' if mode=='genel' else 'Gider_Aktarim']
      self.assertEqual(ws.max_row,1);self.assertEqual(ws.cell(1,3 if mode=='genel' else 5).value,'KEBAPÇI GÖKHAN')
     finally:w.close()
    finally:s.db.close()
 def test_other_errors_still_block_and_identify_record(self):
  d=sample();d.update(toplam='999',aciklama='KEBAPÇI GÖKHAN',vkn='16171222776')
  with self.assertRaisesRegex(ValueError,'Gider kaydı #18.*KEBAPÇI GÖKHAN.*toplamı uyuşmuyor'):validate_export(d,18)
