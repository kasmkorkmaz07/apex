import copy
from pathlib import Path
import sqlite3
import tempfile
import unittest
from core import Store, draft, gross_line, validate, export_xlsx, HEADERS, money

def sample(code='MS', no='000012'):
    d = draft('DENEME FIRMA\nVKN: 0123456789\n13.09.2026\nFIS NO: '+no+'\nKDV %20\nTOPLAM: 120,00')
    d['kod'] = code
    return d

class Tests(unittest.TestCase):
    def test_tax_directory_neighbors_and_not_phone(self):
        self.assertEqual(draft('TEST FIRMA\nKIZILBEY / 1234567890\n01.09.2026')['vkn'],'1234567890')
        self.assertEqual(draft('TEST FIRMA\nKIZILBEY\n1234567890\n01.09.2026')['vkn'],'1234567890')
        self.assertEqual(draft('TEST FIRMA\nKIZILBEY MAH TEL 1234567890')['vkn'],'')

    def test_delete_blank_and_approved_receipts_keeps_audit(self):
        with tempfile.TemporaryDirectory() as tmp:
            s=Store(Path(tmp)/'db');c=s.client('Test');blank=s.add(c,None,draft(),'09-2026')
            r=s.add(c,None,sample(),'09-2026');s.save(r,s.get(r)[1],True)
            self.assertEqual(s.delete_receipts(c,[blank,r]),{'09-2026'})
            self.assertEqual(s.rows(c),[])
            self.assertEqual(s.db.execute('SELECT count(*) FROM deleted_receipts').fetchone()[0],2)
            s.add_vendor('1234567890','CARİ TEST');s.update_vendor('1234567890','12345678901','CARİ YENİ')
            self.assertEqual(s.vendor_rows('yeni')[0][0],'12345678901');s.db.close()

    def test_printed_vat_blocks_wrong_780(self):
        d=sample();d['toplam']='780';d['kalemler']=[gross_line('20','780')];d['ocr']='%20\nTOPLAM *780,00\nTOPKDV *116,67'
        with self.assertRaisesRegex(ValueError,'basılı'):validate(d)

    def test_bim_vdm_consumer_id_and_vat_summary(self):
        raw='BIM BIRLESIK MAGAZALAR A.S.\nBuyuk Mukellefler VDM 1750051846\n02.09.2026\nFATURA NO:T092026007980787\nTCKN/VKN:11111111111-NIHAI TUKETICI\nOdenecek KDV Dahil Tutar *284.00\nKDV MATRAH KDV TUTAR KDV DAHIL\n%1. 167.33 *1.67 *169.00\n%10 104.55 *10.45 *115.00'
        d=draft(raw);validate(d)
        self.assertEqual(d['vkn'],'1750051846')
        self.assertEqual([money(r['matrah'])+money(r['kdv']) for r in d['kalemler']],[money('169'),money('115')])
        self.assertEqual(draft(raw.replace('VDM 1750051846','VDM\n1750051846'))['vkn'],'1750051846')

    def test_background_percentage_and_manual_override(self):
        raw='%10 %3\nA101 YENI MAGAZACILIK A.S.\nUSKUDAR/9480423762\n20.08.2026\nBELGE NO:1\nBOYAMA SETI %10. *159.00\nODENECEK TUTAR 159.00'
        d=draft(raw);validate(d);self.assertEqual(d['kalemler'][0]['oran'],'10.00')
        d['kalemler']=[gross_line('20','159')]
        with self.assertRaises(ValueError):validate(d)
        d['vat_manually_reviewed']=True;validate(d)

    def test_management_search_and_client_archive(self):
        with tempfile.TemporaryDirectory() as tmp:
            s=Store(Path(tmp)/'db');c=s.client('Test')
            s.update_vendor('1234567890','1234567890','A101 YENİ MAĞAZACILIK A.Ş.')
            self.assertEqual(len(s.vendor_rows('yeni magazacilik')),1)
            self.assertEqual(len(s.vendor_rows('a')),1)
            s.rename_client(c,'Yeni');s.archive_client(c)
            self.assertEqual(s.clients(),[]);self.assertEqual(s.clients(True),[(c,'Yeni')])
            s.archive_client(c,False);self.assertEqual(s.clients(),[(c,'Yeni')]);s.db.close()

    def test_posting_month_preserves_document_date_and_updates_same_file(self):
        from core import posting_period
        from openpyxl import load_workbook
        for date,period,expected in [('15.08.2026','09-2026','01.09.2026'),('31.12.2025','08-2026','01.08.2026'),('09.09.2026','09-2026','09.09.2026'),('01.10.2026','09-2026','01.10.2026')]:
            d=posting_period(dict(evrak_tarihi=date,tarih=date),period)
            self.assertEqual(d['tarih'],expected);self.assertEqual(d['evrak_tarihi'],date)
        with tempfile.TemporaryDirectory() as tmp:
            s=Store(Path(tmp)/'db');c=s.client('Test');d=sample();d['evrak_tarihi']='15.08.2026'
            rid=s.add(c,None,d,'09-2026');s.save(rid,s.get(rid)[1],True)
            p=Path(tmp)/'same.xlsx';s.sync_month(c,'09-2026',p)
            d['evrak_no']='99';rid=s.add(c,None,d,'09-2026');s.save(rid,s.get(rid)[1],True)
            self.assertEqual(s.sync_month(c,'09-2026'),p)
            self.assertEqual(s.month_path(c,'09-2026',Path(tmp)/'different.xlsx'),p)
            w=load_workbook(p);self.assertEqual(w['Gider_Aktarim'].max_row,2)
            self.assertEqual(w['Gider_Aktarim']['B1'].value.strftime('%d.%m.%Y'),'01.09.2026')
            w.close();s.db.close()

    def test_vendor_memory_requires_approval_and_survives_restart(self):
        with tempfile.TemporaryDirectory() as tmp:
            p=Path(tmp)/'db.sqlite'; s=Store(p); c=s.client('A')
            d=sample(); rid=s.add(c,None,d)
            d['aciklama']='DOĞRU ÜNVAN'; s.save(rid,d,False)
            self.assertEqual(s.recall_vendor(sample())['aciklama'],'DENEME FIRMA')
            s.save(rid,d,True); s.db.close(); s=Store(p)
            self.assertEqual(s.recall_vendor(sample())['aciklama'],'DOĞRU ÜNVAN')
            s.db.close()

    def test_long_document_and_vat_rows_excluding_zero(self):
        from openpyxl import load_workbook
        raw='A101 YENI MAGAZACILIK A.S.\nUSKUDAR/9480423762\n07.09.2026\nBELGE NO:13029020260907171243\nODENECEK TUTAR *1060,00\nKDV Orani KDV Dahil Tutar KDV\n%20 *500,00\n%10 *250,00\n%1 *300,00\n%0 *10,00'
        d=draft(raw); validate(d)
        self.assertEqual(d['evrak_no'],'9020260907171243')
        self.assertEqual(d['original_evrak_no'],'13029020260907171243')
        self.assertEqual(len(d['kalemler']),3)
        with tempfile.TemporaryDirectory() as tmp:
            p=Path(tmp)/'out.xlsx'; export_xlsx(p,[(1,d)],headers=False)
            w=load_workbook(p); rows=list(w.active.values)
            self.assertEqual([r[5:] for r in rows],[(20,500),(10,250),(1,300)])
            self.assertTrue(all(r[:5]==rows[0][:5] for r in rows)); w.close()

    def test_seller_vd_without_colon_and_date_next_to_time(self):
        for label in ('V.D.','V. D.', 'VD:'):
            for customer in ('MÜŞTERİ', 'MOSTERI', 'ALICI'):
                raw=f'MIGROS TICARET A.S.\nBUYUK MUKELLEFLER {label}6220529513\nTARIH:09/09/2026 SAAT : 20:50\nFIS NO:0263\n{customer} TCKN:11111111111\nSEKER %1 *638,95\nTOPLAM *638,95'
                d=draft(raw)
                self.assertEqual(d['vkn'],'6220529513')
                self.assertEqual(d['evrak_tarihi'],'09.09.2026')
                self.assertEqual(d['evrak_no'],'263')
                self.assertNotIn('11111111111',d['ocr'])
                validate(d)

    def test_customer_identity_never_substitutes_missing_seller(self):
        d=draft('MIGROS TICARET A.S.\nMÜŞTERİ TCKN:11111111111\n09/09/2026\nFIS NO:263')
        self.assertEqual(d['vkn'],'')

    def test_verified_vendor_name_is_bound_to_tax_id(self):
        raw='CGS_CETINKAYA GIYIM SAN. VE TIC.A.S\nVKN: 2470011603\n27.08.2026\nFIS NO: 133\nTOPLAM 179,90\n%20'
        self.assertEqual(draft(raw)['aciklama'],'ÇGS ÇETİNKAYA GİYİM SAN. VE TİC. A.Ş')
        self.assertNotEqual(draft(raw.replace('2470011603','1234567890'))['aciklama'],'ÇGS ÇETİNKAYA GİYİM SAN. VE TİC. A.Ş')

    def test_scanned_discount_receipt_and_number_priority(self):
        d=draft('ÇGS ÇETİNKAYA GİYİM SAN VE TİC A.Ş\nVD:Buyukmukellefler VN:2470011603\n27.08.2026\nFI.$ N0 : 0133\nHAVLU %20 *219,90\nURUN IND. *40,00-D\nTCIPLAM *179,90\nBELGE NO: 00090032708640300135')
        validate(d)
        self.assertEqual(d['evrak_no'],'133')
        self.assertEqual(d['vkn'],'2470011603')
        self.assertEqual(d['kalemler'][0]['kdv'],'29.98')
        self.assertEqual(d['toplam'],'179.90')

    def test_multiple_rates_reconcile_to_receipt_total(self):
        base='TEST MARKET\nVKN:0123456789\n13.09.2026\nFIS NO:1\nURUN %10 *110,00\nURUN %20 *120,00\nTOPLAM *230,00'
        self.assertEqual(len(draft(base)['kalemler']),2)
        self.assertEqual(draft(base.replace('230,00','220,00'))['kalemler'],[])
        d=draft(base+'\nKDV Orani KDV Dahil Tutar KDV\n%10 *110,00 *10,00\n%20 *120,00 *20,00')
        validate(d)
        self.assertEqual(len(d['kalemler']),2)

    def test_person_receipt_and_no_cross_line_number(self):
        text='CTRI\nACR AYAKKABI DUNYASI\nVELİT ACAR\nBAHÇEŞEHİR MAH MEHMET FUAT\nSEYHAN TC: 11872358656\n15-08-2026 FI$ NO: 1\nTERLIK %10 *1.500,00\nTOPLAM *1.500,00'
        d=draft(text)
        validate(d)
        self.assertEqual(d['vkn'],'11872358656')
        self.assertEqual(d['aciklama'],'VELİT ACAR')
        self.assertEqual(d['evrak_no'],'1')
        self.assertEqual(draft(text.replace('NO: 1','NO:'))['evrak_no'],'')
        self.assertEqual(draft(text.replace('VELİT ACAR\n',''))['aciklama'],'')
        d['aciklama']='ACR AYAKKABI DUNYASI'
        with self.assertRaisesRegex(ValueError,'TCKN bulunan'):
            validate(d)

    def test_monthly_update_sections_and_periods(self):
        from openpyxl import load_workbook
        with tempfile.TemporaryDirectory() as tmp:
            s=Store(Path(tmp)/'db.sqlite')
            c=s.client('A'); other=s.client('B')
            p=Path(tmp)/'08-2026.xlsx'
            s.sync_month(c,'08-2026',p)
            for client,period,kind,no in [(c,'08-2026','gider','1'),(c,'08-2026','gelir','2'),(c,'09-2026','gider','3'),(other,'08-2026','gider','4')]:
                d=sample('ML' if kind=='gelir' else 'MS',no)
                d['kind']=kind
                if kind=='gelir':d.update(tarih='02.08.2026',evrak_tarihi='02.08.2026',vkn='')
                rid=s.add(client,'',d,period)
                s.save(rid,d,True)
            s.sync_month(c,'08-2026')
            d=sample(no='5'); rid=s.add(c,'',d,'08-2026'); s.save(rid,d,True)
            s.sync_month(c,'08-2026'); s.sync_month(c,'08-2026')
            w=load_workbook(p)
            self.assertEqual(w.sheetnames,['Giderler','Gelirler','Gider_Aktarim','Gelir_Aktarim'])
            self.assertEqual(w['Giderler']['E1'].value,'Gider Açıklama')
            self.assertEqual(w['Gelirler']['F1'].value,'Satılan Mal/Hizmet')
            self.assertEqual(w['Gider_Aktarim'].max_row,2)
            self.assertEqual(w['Gelir_Aktarim'].max_row,1)
            self.assertEqual(w['Gider_Aktarim']['A1'].value,'MS')
            self.assertEqual(w['Gelir_Aktarim']['A1'].value,'ML')
            w.close()
            self.assertEqual(len(s.rows(c,'gelir')),1)
            self.assertEqual(s.receipt_period(rid),'08-2026')
            s.db.close()
    def test_real_receipt_ocr_regression(self):
        d = draft('EMRE OCAKBA$I\nSeyhan Vd:7351963208\nTARTH: 31-07-2026\nFI$ NO: 0006\nYiyecek %10 *5.900,00\nTOPKDV *536,36\nTOPLAM *5.900.00\nISYERI NO: 001700000020925\nTERMINAL:P2667454\nREF NO:123\nSAAT:01:13\nONAY KODU:996478\nTOPLAM 5.900,00 TL\nTOPLAM 701,23 TL PARAPUANINIZ VAR.')
        validate(d)
        self.assertEqual(d['aciklama'],'EMRE OCAKBAŞI')
        self.assertEqual(d['vkn'],'7351963208')
        self.assertEqual(d['evrak_no'],'6')
        self.assertEqual(d['toplam'],'5900.00')
        self.assertEqual(d['kalemler'],[{'oran':'10.00','matrah':'5363.64','kdv':'536.36'}])
        for word in ('TERMINAL','ISYERI','REF NO','SAAT','ONAY','PARAPUAN'):
            self.assertNotIn(word,d['ocr'])

    def test_parser_and_money(self):
        d = sample()
        validate(d)
        self.assertEqual(d['vkn'],'0123456789')
        self.assertEqual(d['evrak_no'],'12')
        self.assertEqual(money('1.234,56'),money('1234.56'))
        self.assertEqual(d['kalemler'][0]['matrah'],'100.00')
        self.assertEqual(draft('TOPLAM 120,00\n%10\n%20')['kalemler'],[])
        self.assertEqual(draft('VKN: 0123456789\nVKN: 1123456789')['vkn'],'')

    def test_missing_and_mismatch(self):
        d = sample()
        d['toplam'] = '125'
        with self.assertRaises(ValueError): validate(d)
        d = sample()
        d['evrak_tarihi'] = '31.02.2026'
        with self.assertRaises(ValueError): validate(d)
        with self.assertRaises(ValueError): money('NaN')
        with self.assertRaises(ValueError): gross_line('20','-1')

    def test_wrong_vat_rate_cannot_be_approved(self):
        d = draft('EMRE OCAKBA$I\nSeyhan Vd:7351963208\n31-07-2026\nFI$ NO: 0006\nYiyecek %10 *5.900,00\nTOPLAM *5.900.00')
        validate(d)
        d['kalemler'] = [gross_line('20','5900')]
        with self.assertRaisesRegex(ValueError,'Fişte okunan KDV oranı'):
            validate(d)

    def test_store_isolation_duplicate_export(self):
        with tempfile.TemporaryDirectory() as tmp:
            s = Store(Path(tmp)/'test.db')
            a,b = s.client('A'),s.client('B')
            i = s.add(a,'',sample())
            s.save(i,sample(),True)
            j = s.add(a,'',sample('ML','000013'))
            s.save(j,sample('ML','000013'),True)
            k = s.add(a,'',sample())
            with self.assertRaises(sqlite3.IntegrityError): s.save(k,sample(),True)
            other = s.add(b,'',sample())
            s.save(other,sample(),True)
            self.assertEqual(len(s.pending(a,'09-2026')),2)
            self.assertEqual(len(s.pending(a,'08-2026')),0)
            s.mark_exported([i],'sample.xlsx')
            self.assertEqual(len(s.pending(a,'09-2026')),1)
            with self.assertRaises(ValueError): s.save(i,sample())
            s.db.close()

    def test_seven_columns_mixed_codes_and_gross(self):
        from openpyxl import load_workbook
        with tempfile.TemporaryDirectory() as tmp:
            a,b = sample(),sample('ML','000013')
            b['kalemler'] = [gross_line('10','110'),gross_line('20','120')]
            b['ocr'] = b['ocr'] + '\nKDV %10'
            b['toplam'] = '230'
            b['aciklama'] = '=NOT_A_FORMULA'
            p = Path(tmp)/'export.xlsx'
            export_xlsx(p,[(1,a),(2,b)])
            w = load_workbook(p)
            s = w['Isletme']
            self.assertEqual(s.max_column,7)
            self.assertEqual([c.value for c in s[1]],HEADERS)
            self.assertEqual([s.cell(r,1).value for r in (2,3,4)],['MS','ML','ML'])
            self.assertEqual([s.cell(r,7).value for r in (2,3,4)],[120,110,120])
            self.assertEqual(s['C2'].value,'12')
            self.assertEqual(s['D2'].value,'0123456789')
            self.assertEqual(s['E3'].data_type,'s')
            self.assertEqual(s['B2'].value.day,13)
            w.close()

if __name__=='__main__':
    unittest.main()

