import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch
from openpyxl import Workbook, load_workbook
from excel_save import save_workbook


class ExcelSaveTests(unittest.TestCase):
    def test_identifies_zirve_lock_without_closing_it(self):
        with tempfile.TemporaryDirectory() as tmp:
            path=Path(tmp)/'mukellef.xlsx'
            with patch.object(Path,'replace',side_effect=PermissionError()), patch('file_lock.owners',return_value=['ZirveMusavir.exe']):
                with self.assertRaisesRegex(PermissionError,'ZirveMusavir.exe'):
                    save_workbook(Workbook(),path)

    def test_locked_target_preserved_and_retry_updates_same_file(self):
        with tempfile.TemporaryDirectory() as tmp:
            path=Path(tmp)/'mukellef.xlsx'
            wb=Workbook();wb.active['A1']='Eski ünvan';save_workbook(wb,path)
            before=path.read_bytes();wb.active['A1']='Düzeltilmiş ünvan'
            with patch.object(Path,'replace',side_effect=PermissionError('WinError 5')):
                with self.assertRaisesRegex(PermissionError,'Excel’de veya Zirve’de'):
                    save_workbook(wb,path)
            self.assertEqual(path.read_bytes(),before)
            self.assertEqual(list(Path(tmp).iterdir()),[path])
            save_workbook(wb,path)
            result=load_workbook(path);self.assertEqual(result.active['A1'].value,'Düzeltilmiş ünvan');result.close()

    def test_failed_workbook_generation_preserves_original(self):
        with tempfile.TemporaryDirectory() as tmp:
            path=Path(tmp)/'mukellef.xlsx';wb=Workbook();save_workbook(wb,path);before=path.read_bytes()
            with patch.object(wb,'save',side_effect=ValueError('Bad value')):
                with self.assertRaises(ValueError):save_workbook(wb,path)
            self.assertEqual(path.read_bytes(),before)
            self.assertEqual(list(Path(tmp).iterdir()),[path])
