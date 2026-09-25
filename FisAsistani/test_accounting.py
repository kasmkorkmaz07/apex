import unittest,tempfile
from pathlib import Path
from copy import deepcopy
from openpyxl import load_workbook
from core import Store,gross_line
from accounting import EMIR,profile,save_profile
class AccountingTests(unittest.TestCase):
 def test_monthly_general_and_business_isolation(self):
  with tempfile.TemporaryDirectory() as tmp:
   s=Store(Path(tmp)/'db');cid=s.client('Emir');other=s.client('Isletme')
   save_profile(s,cid,deepcopy(EMIR));self.assertEqual(profile(s,other)['mode'],'isletme')
   for rate in (1,10,20):
    d=dict(tarih='14.09.2026',evrak_tarihi='14.09.2026',vkn='1234567890',aciklama='Test Firma',evrak_no=str(rate),kod='MS',toplam=str(100+rate),kalemler=[gross_line(rate,100+rate)],kind='gider',ocr='')
    rid=s.add(cid,None,d,'09-2026');s.save(rid,d,True)
   p=Path(tmp)/'out.xlsx';s.sync_month(cid,'09-2026',p);s.sync_month(cid,'09-2026',p)
   w=load_workbook(p);raw=w['Genel_Alis_Aktarim'];self.assertEqual(raw.max_row,3)
   for row in raw.iter_rows(values_only=True):
    rate=int(row[4]);self.assertEqual(row[5]+row[6],100+rate);self.assertEqual(list(row[7:9]),EMIR['rates'][str(rate)]);self.assertEqual(row[9],'320 01')
   w.close();s.db.close()
